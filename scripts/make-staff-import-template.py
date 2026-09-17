#!/usr/bin/env python3
"""Generate the staff bulk-upload workbook handed to clubs.

The counterpart to make-rider-import-template.py. Same reasoning applies: the
columns must track rowSchema in app/api/staff/import/route.ts, and a template
that drifts from the importer is worse than none.

    pip3 install openpyxl && python3 scripts/make-staff-import-template.py

Three decisions worth knowing:

* Every column is formatted as TEXT, for the same reason as the rider sheet —
  Excel turns 2026-04-01 into a date serial and re-emits it as 01-04-2026, and
  strips the leading zero off a mobile.
* No cell comments. openpyxl writes them in a form ExcelJS cannot read, and the
  importer parses this workbook directly.
* NO PASSWORD COLUMN, and this one is deliberate rather than an omission. A
  spreadsheet of plaintext staff passwords gets emailed around and left in
  Downloads. The importer generates one per account instead and keeps it
  encrypted for the Credential Sheet to reprint.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT = "public/templates/equiwings-staff-import-template.xlsx"

# Must match ASSIGNABLE_STAFF_ROLES in lib/schemas/staff.ts — i.e. ROLES minus
# SUPER_ADMIN, ADMIN, RIDER, PARENT and EXAMINER.
ROLES = [
    "CENTRE_MANAGER", "HEAD_COACH", "COACH", "STABLE_MANAGER", "INVENTORY_MANAGER",
    "GROOM", "FARRIER", "VET", "ACCOUNTANT", "SCHOOL_ADMINISTRATOR",
]

# (header, required, width, help) — order matches the CSV the importer expects.
COLUMNS = [
    ("name",         True,  24, "Full name of the employee. Required."),
    ("email",        True,  30, "REQUIRED — this IS their login.\n"
                                "Must be unique across the whole platform, not just this centre.\n"
                                "A staff row without one creates an account nobody can sign into."),
    ("role",         True,  22, "REQUIRED. Pick from the dropdown, or type a normal job title —\n"
                                "'Head Coach', 'Instructor', 'Syce', 'Storekeeper', 'Accounts' are\n"
                                "all understood. An unrecognised title is reported, never guessed."),
    ("phone",        False, 16, "Optional. 10-digit Indian mobile — 9876543210.\n"
                                "+91, spaces and a 0-prefix are accepted and cleaned.\n"
                                "This is the fallback for password reset, so a wrong number costs\n"
                                "somebody their login later."),
    ("salary_band",  False, 14, "Optional. Your own grade code — C1, C2, A3. Free text."),
    ("joining_date", False, 14, "Optional. YYYY-MM-DD. Defaults to the upload date if blank.\n"
                                "Set it for people who joined the club before it went on the system."),
]

HEAD_REQ = PatternFill("solid", fgColor="1F3864")   # deep blue: required
HEAD_OPT = PatternFill("solid", fgColor="4A6FA5")   # lighter: optional
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)

wb = Workbook()

# ── Sheet 1: the data the club fills in ──────────────────────────────────────
ws = wb.active
# The importer reads the FIRST sheet, but name it for the human either way.
ws.title = "Staff"

for i, (name, required, width, help_text) in enumerate(COLUMNS, start=1):
    letter = get_column_letter(i)
    c = ws.cell(row=1, column=i, value=name)
    c.font = Font(bold=True, color="FFFFFF", size=11)
    c.fill = HEAD_REQ if required else HEAD_OPT
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = THIN
    ws.column_dimensions[letter].width = width
    for r in range(2, 1002):
        ws.cell(row=r, column=i).number_format = "@"

ws.row_dimensions[1].height = 30
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"

# Role dropdown — the one column where a typo is both likely and consequential,
# since the role decides what this person can see.
role_col = get_column_letter([c[0] for c in COLUMNS].index("role") + 1)
dv = DataValidation(
    type="list",
    formula1='"' + ",".join(ROLES) + '"',
    allow_blank=False,
    showErrorMessage=False,  # warn-only: typed job titles are resolved server-side
)
ws.add_data_validation(dv)
dv.add(f"{role_col}2:{role_col}1001")

# ── Sheet 2: how to use it ───────────────────────────────────────────────────
info = wb.create_sheet("Instructions")
info.column_dimensions["A"].width = 4
info.column_dimensions["B"].width = 104

LINES = [
    ("h", "Equiwings — bulk staff upload"),
    ("", ""),
    ("b", "One file per centre"),
    ("p", "Staff are created in the centre selected in the top bar when you upload. There is no"),
    ("p", "'centre' column — the centre comes from where you upload, not from the file. If you"),
    ("p", "run more than one centre, keep a separate copy of this file per centre."),
    ("", ""),
    ("b", "Steps"),
    ("p", "1.  Fill in the 'Staff' sheet. One person per row. Do not rename or reorder columns."),
    ("p", "2.  In Equiwings: Staff > Bulk Upload Staff, and choose this .xlsx file directly."),
    ("p", "3.  Press Preview first. It reports duplicate emails and unknown roles WITHOUT"),
    ("p", "    saving anything. Fix those rows, then press Create."),
    ("", ""),
    ("b", "There is no password column, on purpose"),
    ("p", "Each account gets its own randomly generated password. It is stored encrypted and"),
    ("p", "can be reprinted from Users > Credential Sheet when you hand the login over."),
    ("p", "Putting passwords in a spreadsheet means emailing them around and leaving them in"),
    ("p", "everyone's Downloads folder, which is exactly what this avoids."),
    ("", ""),
    ("b", "Email is the login, and it must be unique"),
    ("p", "One address, one person, across the entire platform — not just this centre. If an"),
    ("p", "address is already in use the row is skipped and reported; nothing is overwritten."),
    ("p", "Shared inboxes (office@, accounts@) can only ever belong to one account."),
    ("", ""),
    ("b", "Roles"),
    ("p", "The role column decides what this person can see, so it is the one cell worth"),
    ("p", "checking twice. Pick from the dropdown, or type the job title you normally use —"),
    ("p", "'Head Coach', 'Instructor', 'Syce', 'Storekeeper', 'Accounts' are all understood."),
    ("p", "Anything unrecognised is reported by name rather than guessed at."),
    ("m", "  " + "  ".join(ROLES[:5])),
    ("m", "  " + "  ".join(ROLES[5:])),
    ("", ""),
    ("b", "Dates"),
    ("p", "Type joining_date as plain text in the form 2026-04-01. The columns in this file are"),
    ("p", "pre-set to Text to stop Excel rewriting them — if you paste from elsewhere, use"),
    ("p", "Paste Special > Values."),
    ("", ""),
    ("b", "Example rows (do not paste these in — they are here so the sheet stays clean)"),
    ("m", "name           email            role        phone       salary_band  joining_date"),
    ("m", "Ravi Kumar     ravi@club.in     COACH       9876543210  C2           2026-04-01"),
    ("m", "Meena Rao      meena@club.in    Head Coach  9876501234  C1           2026-04-01"),
    ("m", "Imran Shaikh   imran@club.in    Syce        9876512345               2026-05-15"),
    ("", ""),
    ("b", "What this does NOT create"),
    ("p", "KYC documents. Aadhaar, PAN, bank proof and police verification cannot travel in a"),
    ("p", "spreadsheet. Send new hires the self-onboarding link from Staff > Employee"),
    ("p", "Onboarding so they upload their own, or attach them per person afterwards."),
    ("p", "Police verification is mandatory before anyone works with minors, and a bulk upload"),
    ("p", "must not become a way to skip it."),
]
r = 2
for kind, text in LINES:
    cell = info.cell(row=r, column=2, value=text)
    if kind == "h":
        cell.font = Font(bold=True, size=16, color="1F3864")
        info.row_dimensions[r].height = 24
    elif kind == "b":
        cell.font = Font(bold=True, size=11, color="1F3864")
    elif kind == "m":
        cell.font = Font(name="Menlo", size=9, color="444444")
    else:
        cell.font = Font(size=11)
    r += 1

wb.save(OUT)
print(f"wrote {OUT}")
